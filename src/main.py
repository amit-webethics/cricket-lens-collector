import asyncio
import hashlib
import io
import json
import re
from datetime import datetime, timedelta, timezone
from html import unescape
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode

import feedparser
import httpx
from apify import Actor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

HEADERS = ['Story ID', 'Source', 'Original headline', 'Source URL', 'Published UTC',
           'Collected UTC', 'Suggested type', 'Feed summary (source text)', 'Author',
           'Source tags', 'Date check', 'Review status', 'Your headline', 'Your article',
           'Series / category', 'Image reference / licence', 'Editorial notes']

def clean(value):
    return ILLEGAL_CHARACTERS_RE.sub('', unescape(re.sub(r'<[^>]+>', ' ', str(value or '')))).strip()

def canonical(url):
    p = urlsplit(url)
    if p.scheme not in ('http', 'https') or not p.hostname:
        return ''
    query = [(k, v) for k, v in parse_qsl(p.query) if not k.lower().startswith('utm_') and k.lower() not in ('fbclid', 'gclid')]
    return urlunsplit((p.scheme.lower(), p.netloc.lower(), p.path.rstrip('/'), urlencode(sorted(query)), ''))

def classify(title):
    title = title.lower()
    for label, words in [('Preview', ('preview', 'prediction')), ('Opinion / analysis', ('opinion', 'analysis', 'column')), ('Review / report', ('report', 'review', 'highlights'))]:
        if any(re.search(r'\b' + word + r'\b', title) for word in words):
            return label
    return 'News'

def make_row(entry, source, now, cutoff):
    url = canonical(entry.get('link', ''))
    title = clean(entry.get('title'))
    if not url or not title:
        return None
    parsed = entry.get('published_parsed') or entry.get('updated_parsed')
    stamp = None
    if parsed:
        try:
            stamp = datetime(*parsed[:6], tzinfo=timezone.utc)
        except (ValueError, TypeError):
            pass
    if stamp and stamp < cutoff:
        return None
    return dict(zip(HEADERS, [hashlib.sha256(url.encode()).hexdigest()[:20], source, title, url,
        stamp.isoformat() if stamp else '', now.isoformat(), classify(title),
        clean(entry.get('summary', ''))[:2000], clean(entry.get('author')),
        ', '.join(clean(t.get('term')) for t in entry.get('tags', [])),
        'Check future date' if stamp and stamp > now + timedelta(minutes=10) else ('Available' if stamp else 'Missing — check source'),
        'New', '', '', '', '', '']))

async def collect(config):
    now = datetime.now(timezone.utc)
    hours = max(1, min(720, int(config.get('lookbackHours', 72))))
    limit = max(1, min(500, int(config.get('maxItemsPerSource', 100))))
    sources = config.get('sources', [])
    if not isinstance(sources, list) or not sources:
        raise ValueError('Provide at least one RSS / Atom source.')
    rows, reports, seen = [], [], set()
    async with httpx.AsyncClient(timeout=30, follow_redirects=True, headers={'User-Agent':'CricketLensFeedCollector/0.1'}) as client:
        for source in sources:
            name = clean(source.get('name', 'Unnamed source'))
            url = source.get('url', '')
            report = {'Source': name, 'Feed URL': url, 'Status': '', 'Stories added': 0, 'Detail': ''}
            try:
                if urlsplit(url).scheme != 'https':
                    raise ValueError('An HTTPS feed URL is required')
                response = await client.get(url)
                response.raise_for_status()
                if len(response.content) > 10_000_000:
                    raise ValueError('Feed exceeds 10 MB limit')
                feed = feedparser.parse(response.content)
                if not feed.get('version'):
                    raise ValueError('Response is not a recognised RSS / Atom feed')
                for entry in feed.entries[:limit]:
                    row = make_row(entry, name, now, now - timedelta(hours=hours))
                    if row and row['Source URL'] not in seen:
                        seen.add(row['Source URL'])
                        rows.append(row)
                        report['Stories added'] += 1
                report['Status'] = 'Warning' if feed.bozo else 'OK'
                report['Detail'] = 'Feed has parse warnings; check results' if feed.bozo else ''
            except Exception as exc:
                report['Status'] = 'Failed'
                report['Detail'] = str(exc)[:400]
            reports.append(report)
    rows.sort(key=lambda row: row['Published UTC'], reverse=True)
    return rows, reports

def workbook(rows, reports):
    wb = Workbook()
    wb.remove(wb.active)
    for name, records, headers in [('Stories', rows, HEADERS), ('Source status', reports, ['Source','Feed URL','Status','Stories added','Detail'])]:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for record in records:
            ws.append([record.get(h, '') for h in headers])
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Untrusted feed text must never become an Excel formula.
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='33319C')
            cell.font = Font(color='FFFFFF', bold=True)
        ws.freeze_panes = 'D2' if name == 'Stories' else 'A2'
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = 24
        if name == 'Stories':
            for col in ('C', 'D', 'H', 'N', 'Q'):
                ws.column_dimensions[col].width = 55
            validation = DataValidation(type='list', formula1='"New,Reviewing,Ready,Published,Rejected"')
            ws.add_data_validation(validation)
            validation.add('L2:L10000')
    notes = wb.create_sheet('Instructions')
    for line in ['The Cricket Lens — collection workbook',
                 'Source headlines and summaries are reference material, not original drafts.',
                 'Review dates, facts, suggested type, and image rights before publishing.',
                 'Write your content in Your headline / Your article; record your review status.',
                 'Each run creates a fresh workbook. Keep your edited workbook separately.',
                 'Duplicates are removed by canonical URL within a run only.',
                 'Different sources covering the same event remain separate rows.',
                 'Coverage is limited to configured feeds; missing dates are flagged.',
                 'Source status records failures and parse warnings. No full article scraping or AI rewriting.']:
        notes.append([line])
    notes.column_dimensions['A'].width = 110
    data = io.BytesIO()
    wb.save(data)
    return data.getvalue()

async def main():
    async with Actor:
        config = await Actor.get_input() or json.loads(Path('INPUT.json').read_text())
        rows, reports = await collect(config)
        await Actor.set_value('OUTPUT.xlsx', workbook(rows, reports), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        await Actor.set_value('SOURCE_STATUS', reports)
        if rows:
            await Actor.push_data(rows)
        Actor.log.info(f'Collected {len(rows)} stories. Download OUTPUT.xlsx from run Storage / Key-value store.')
        if all(r['Status'] == 'Failed' for r in reports):
            raise RuntimeError('All feeds failed. See SOURCE_STATUS and OUTPUT.xlsx for details.')

if __name__ == '__main__':
    asyncio.run(main())
